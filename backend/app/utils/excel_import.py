from io import BytesIO
from typing import Any
from openpyxl import load_workbook


HEADER_MAP = {
    "名称": "name",
    "节点名称": "name",
    "景点名称": "name",
    "name": "name",
    "描述": "description",
    "简介": "description",
    "description": "description",
    "地址": "address",
    "address": "address",
    "纬度": "latitude",
    "latitude": "latitude",
    "lat": "latitude",
    "经度": "longitude",
    "longitude": "longitude",
    "lng": "longitude",
    "图片": "image_url",
    "图片地址": "image_url",
    "image_url": "image_url",
    "标签": "tags",
    "类型": "tags",
    "tags": "tags",
}


def parse_points_excel(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [HEADER_MAP.get(str(cell).strip(), str(cell).strip()) if cell else "" for cell in rows[0]]
    points = []
    for row in rows[1:]:
        item = {}
        for index, value in enumerate(row):
            if index >= len(headers) or not headers[index]:
                continue
            key = headers[index]
            if key in {"name", "description", "address", "latitude", "longitude", "image_url", "tags"}:
                item[key] = value
        if any(value not in (None, "") for value in item.values()):
            points.append(item)
    return points
